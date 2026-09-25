import json
import re
from datetime import datetime
from io import BytesIO

from bson import ObjectId
from flask import (
    Blueprint, render_template, request, redirect,
    url_for, session, flash, send_file, Response
)

from openpyxl import Workbook, load_workbook
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from pypdf import PdfReader

import extensions as ext


client_bp = Blueprint("clients", __name__)


def db():
    return ext.db


def login_required():
    return session.get("user_id") and session.get("company_id")


def company_id():
    return session.get("company_id")


def clean_client_doc(client):
    return {
        "name": str(client.get("name", "")).strip(),
        "email": str(client.get("email", "")).strip(),
        "phone": str(client.get("phone", "")).strip(),
        "address": str(client.get("address", "")).strip(),
        "notes": str(client.get("notes", "")).strip(),
    }


def build_client_doc(data):
    clean = clean_client_doc(data)

    return {
        "company_id": company_id(),
        "company_name": session.get("company_name", ""),
        "name": clean["name"],
        "email": clean["email"],
        "phone": clean["phone"],
        "address": clean["address"],
        "notes": clean["notes"],
        "is_active": True,
        "created_by": session.get("user_id"),
        "created_by_name": session.get("user_name", ""),
        "created_at": datetime.utcnow(),
        "updated_at": datetime.utcnow()
    }


def serialize_client(client):
    return {
        "name": client.get("name", ""),
        "email": client.get("email", ""),
        "phone": client.get("phone", ""),
        "address": client.get("address", ""),
        "notes": client.get("notes", ""),
    }


def get_company_clients():
    clients = list(
        db()["clients"]
        .find({"company_id": company_id()})
        .sort("name", 1)
    )

    for c in clients:
        c["_id"] = str(c["_id"])

    return clients


def parse_sql_values(sql_text):
    """
    Espera INSERTS tipo:
    INSERT INTO clients (name,email,phone,address,notes) VALUES ('Juan','a@b.com','555','Dirección','Nota');
    """
    clients = []

    pattern = re.compile(
        r"INSERT\s+INTO\s+clients\s*\((.*?)\)\s*VALUES\s*\((.*?)\)\s*;",
        re.IGNORECASE | re.DOTALL
    )

    for match in pattern.finditer(sql_text):
        columns_raw = match.group(1)
        values_raw = match.group(2)

        columns = [c.strip().replace("`", "").replace('"', "") for c in columns_raw.split(",")]

        values = re.findall(r"'((?:\\'|[^'])*)'|([^,]+)", values_raw)
        clean_values = []

        for quoted, unquoted in values:
            value = quoted if quoted else unquoted
            value = value.strip()
            value = value.replace("\\'", "'")
            clean_values.append(value)

        data = dict(zip(columns, clean_values))

        clients.append({
            "name": data.get("name", ""),
            "email": data.get("email", ""),
            "phone": data.get("phone", ""),
            "address": data.get("address", ""),
            "notes": data.get("notes", "")
        })

    return clients


def parse_pdf_clients(file_storage):
    """
    Importación básica desde PDF.
    Lee líneas con separador | en este orden:
    Nombre | Correo | Teléfono | Dirección | Notas
    """
    reader = PdfReader(file_storage)
    text = ""

    for page in reader.pages:
        text += page.extract_text() or ""

    clients = []

    for line in text.splitlines():
        line = line.strip()

        if not line or "|" not in line:
            continue

        parts = [p.strip() for p in line.split("|")]

        if len(parts) >= 4 and parts[0].lower() not in ["nombre", "cliente"]:
            clients.append({
                "name": parts[0],
                "email": parts[1] if len(parts) > 1 else "",
                "phone": parts[2] if len(parts) > 2 else "",
                "address": parts[3] if len(parts) > 3 else "",
                "notes": parts[4] if len(parts) > 4 else ""
            })

    return clients


def insert_clients_from_list(clients):
    inserted = 0
    duplicated = 0
    errors = 0

    for client in clients:
        clean = clean_client_doc(client)

        if not clean["name"]:
            errors += 1
            continue

        existing = db()["clients"].find_one({
            "company_id": company_id(),
            "name": clean["name"],
            "email": clean["email"]
        })

        if existing:
            duplicated += 1
            continue

        db()["clients"].insert_one(build_client_doc(clean))
        inserted += 1

    return inserted, duplicated, errors


@client_bp.route("/clientes")
def clients_list():
    if not login_required():
        return redirect(url_for("login"))

    q = request.args.get("q", "").strip()

    query = {
        "company_id": company_id()
    }

    if q:
        query["$or"] = [
            {"name": {"$regex": q, "$options": "i"}},
            {"email": {"$regex": q, "$options": "i"}},
            {"phone": {"$regex": q, "$options": "i"}},
            {"address": {"$regex": q, "$options": "i"}},
        ]

    clients = list(
        db()["clients"]
        .find(query)
        .sort("created_at", -1)
    )

    for c in clients:
        c["_id"] = str(c["_id"])

    return render_template(
        "clientes/list.html",
        clients=clients,
        q=q
    )


@client_bp.route("/clientes/add", methods=["GET", "POST"])
def add_client():
    if not login_required():
        return redirect(url_for("login"))

    if request.method == "POST":
        name = request.form.get("name", "").strip()
        email = request.form.get("email", "").strip()
        phone = request.form.get("phone", "").strip()
        address = request.form.get("address", "").strip()
        notes = request.form.get("notes", "").strip()

        if not name:
            flash("El nombre del cliente es obligatorio.", "warning")
            return redirect(url_for("clients.add_client"))

        client_doc = build_client_doc({
            "name": name,
            "email": email,
            "phone": phone,
            "address": address,
            "notes": notes
        })

        db()["clients"].insert_one(client_doc)

        flash("Cliente registrado correctamente ✅", "success")
        return redirect(url_for("clients.clients_list"))

    return render_template("clientes/new.html")


@client_bp.route("/clientes/<client_id>/edit", methods=["GET", "POST"])
def edit_client(client_id):
    if not login_required():
        return redirect(url_for("login"))

    try:
        oid = ObjectId(client_id)
    except Exception:
        flash("Cliente inválido.", "danger")
        return redirect(url_for("clients.clients_list"))

    client = db()["clients"].find_one({
        "_id": oid,
        "company_id": company_id()
    })

    if not client:
        flash("Cliente no encontrado o no pertenece a esta empresa.", "danger")
        return redirect(url_for("clients.clients_list"))

    if request.method == "POST":
        name = request.form.get("name", "").strip()
        email = request.form.get("email", "").strip()
        phone = request.form.get("phone", "").strip()
        address = request.form.get("address", "").strip()
        notes = request.form.get("notes", "").strip()

        if not name:
            flash("El nombre del cliente es obligatorio.", "warning")
            return redirect(url_for("clients.edit_client", client_id=client_id))

        db()["clients"].update_one(
            {
                "_id": oid,
                "company_id": company_id()
            },
            {
                "$set": {
                    "name": name,
                    "email": email,
                    "phone": phone,
                    "address": address,
                    "notes": notes,
                    "updated_at": datetime.utcnow()
                }
            }
        )

        flash("Cliente actualizado correctamente ✅", "success")
        return redirect(url_for("clients.clients_list"))

    client["_id"] = str(client["_id"])

    return render_template(
        "clientes/edit.html",
        client=client
    )


@client_bp.route("/clientes/<client_id>/delete", methods=["POST"])
def delete_client(client_id):
    if not login_required():
        return redirect(url_for("login"))

    try:
        oid = ObjectId(client_id)
    except Exception:
        flash("Cliente inválido.", "danger")
        return redirect(url_for("clients.clients_list"))

    result = db()["clients"].delete_one({
        "_id": oid,
        "company_id": company_id()
    })

    if result.deleted_count == 0:
        flash("No se pudo eliminar el cliente.", "danger")
    else:
        flash("Cliente eliminado correctamente.", "success")

    return redirect(url_for("clients.clients_list"))


@client_bp.route("/clientes/import", methods=["POST"])
def import_clients():
    if not login_required():
        return redirect(url_for("login"))

    file = request.files.get("file")

    if not file or not file.filename:
        flash("Selecciona un archivo para importar.", "warning")
        return redirect(url_for("clients.clients_list"))

    filename = file.filename.lower()
    clients = []

    try:
        if filename.endswith(".json"):
            data = json.load(file)

            if isinstance(data, dict):
                clients = data.get("clients", [])
            elif isinstance(data, list):
                clients = data

        elif filename.endswith(".xlsx"):
            wb = load_workbook(file)
            ws = wb.active

            headers = []
            for cell in ws[1]:
                headers.append(str(cell.value).strip().lower() if cell.value else "")

            for row in ws.iter_rows(min_row=2, values_only=True):
                data = dict(zip(headers, row))

                clients.append({
                    "name": data.get("name") or data.get("nombre") or "",
                    "email": data.get("email") or data.get("correo") or "",
                    "phone": data.get("phone") or data.get("telefono") or data.get("teléfono") or "",
                    "address": data.get("address") or data.get("direccion") or data.get("dirección") or "",
                    "notes": data.get("notes") or data.get("notas") or "",
                })

        elif filename.endswith(".sql"):
            sql_text = file.read().decode("utf-8", errors="ignore")
            clients = parse_sql_values(sql_text)

        elif filename.endswith(".pdf"):
            clients = parse_pdf_clients(file)

        else:
            flash("Formato no permitido. Usa JSON, Excel, SQL o PDF.", "warning")
            return redirect(url_for("clients.clients_list"))

        inserted, duplicated, errors = insert_clients_from_list(clients)

        flash(
            f"Importación finalizada. Insertados: {inserted}, duplicados: {duplicated}, errores: {errors}.",
            "success"
        )

    except Exception as e:
        flash(f"Error al importar clientes: {str(e)}", "danger")

    return redirect(url_for("clients.clients_list"))


@client_bp.route("/clientes/export/<fmt>")
def export_clients(fmt):
    if not login_required():
        return redirect(url_for("login"))

    clients = get_company_clients()
    fmt = fmt.lower()

    if fmt == "json":
        data = [serialize_client(c) for c in clients]

        return Response(
            json.dumps(data, ensure_ascii=False, indent=2),
            mimetype="application/json",
            headers={
                "Content-Disposition": "attachment; filename=clientes.json"
            }
        )

    if fmt == "xlsx":
        wb = Workbook()
        ws = wb.active
        ws.title = "Clientes"

        ws.append(["name", "email", "phone", "address", "notes"])

        for c in clients:
            ws.append([
                c.get("name", ""),
                c.get("email", ""),
                c.get("phone", ""),
                c.get("address", ""),
                c.get("notes", ""),
            ])

        output = BytesIO()
        wb.save(output)
        output.seek(0)

        return send_file(
            output,
            as_attachment=True,
            download_name="clientes.xlsx",
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

    if fmt == "sql":
        lines = []

        for c in clients:
            name = str(c.get("name", "")).replace("'", "\\'")
            email = str(c.get("email", "")).replace("'", "\\'")
            phone = str(c.get("phone", "")).replace("'", "\\'")
            address = str(c.get("address", "")).replace("'", "\\'")
            notes = str(c.get("notes", "")).replace("'", "\\'")

            lines.append(
                "INSERT INTO clients (name,email,phone,address,notes) "
                f"VALUES ('{name}','{email}','{phone}','{address}','{notes}');"
            )

        sql = "\n".join(lines)

        return Response(
            sql,
            mimetype="text/plain",
            headers={
                "Content-Disposition": "attachment; filename=clientes.sql"
            }
        )

    if fmt == "pdf":
        output = BytesIO()
        pdf = canvas.Canvas(output, pagesize=letter)
        width, height = letter

        y = height - 50

        pdf.setFont("Helvetica-Bold", 14)
        pdf.drawString(40, y, "Clientes")
        y -= 20

        pdf.setFont("Helvetica", 9)
        pdf.drawString(40, y, f"Empresa: {session.get('company_name', '')}")
        y -= 30

        pdf.setFont("Helvetica-Bold", 8)
        pdf.drawString(40, y, "Nombre | Correo | Teléfono | Dirección | Notas")
        y -= 15

        pdf.setFont("Helvetica", 8)

        for c in clients:
            line = (
                f"{c.get('name', '')} | "
                f"{c.get('email', '')} | "
                f"{c.get('phone', '')} | "
                f"{c.get('address', '')} | "
                f"{c.get('notes', '')}"
            )

            if len(line) > 120:
                line = line[:117] + "..."

            pdf.drawString(40, y, line)
            y -= 14

            if y < 50:
                pdf.showPage()
                y = height - 50
                pdf.setFont("Helvetica", 8)

        pdf.save()
        output.seek(0)

        return send_file(
            output,
            as_attachment=True,
            download_name="clientes.pdf",
            mimetype="application/pdf"
        )

    flash("Formato de exportación no válido.", "warning")
    return redirect(url_for("clients.clients_list"))


@client_bp.route("/clientes/backups/create", methods=["POST"])
def create_clients_backup():
    if not login_required():
        return redirect(url_for("login"))

    clients = get_company_clients()

    backup_doc = {
        "company_id": company_id(),
        "company_name": session.get("company_name", ""),
        "module": "clients",
        "total_clients": len(clients),
        "clients": [serialize_client(c) for c in clients],
        "created_by": session.get("user_id"),
        "created_by_name": session.get("user_name", ""),
        "created_at": datetime.utcnow()
    }

    db()["client_backups"].insert_one(backup_doc)

    flash("Backup de clientes creado correctamente ✅", "success")
    return redirect(url_for("clients.clients_list"))


@client_bp.route("/clientes/backups")
def clients_backups():
    if not login_required():
        return redirect(url_for("login"))

    backups = list(
        db()["client_backups"]
        .find({"company_id": company_id(), "module": "clients"})
        .sort("created_at", -1)
    )

    for b in backups:
        b["_id"] = str(b["_id"])

    return render_template(
        "clientes/backups.html",
        backups=backups
    )


@client_bp.route("/clientes/backups/<backup_id>/download/<fmt>")
def download_clients_backup(backup_id, fmt):
    if not login_required():
        return redirect(url_for("login"))

    try:
        oid = ObjectId(backup_id)
    except Exception:
        flash("Backup inválido.", "danger")
        return redirect(url_for("clients.clients_backups"))

    backup = db()["client_backups"].find_one({
        "_id": oid,
        "company_id": company_id(),
        "module": "clients"
    })

    if not backup:
        flash("Backup no encontrado.", "danger")
        return redirect(url_for("clients.clients_backups"))

    clients = backup.get("clients", [])
    fmt = fmt.lower()

    filename_base = f"backup_clientes_{backup_id}"

    if fmt == "json":
        return Response(
            json.dumps(clients, ensure_ascii=False, indent=2),
            mimetype="application/json",
            headers={
                "Content-Disposition": f"attachment; filename={filename_base}.json"
            }
        )

    if fmt == "xlsx":
        wb = Workbook()
        ws = wb.active
        ws.title = "Backup Clientes"

        ws.append(["name", "email", "phone", "address", "notes"])

        for c in clients:
            ws.append([
                c.get("name", ""),
                c.get("email", ""),
                c.get("phone", ""),
                c.get("address", ""),
                c.get("notes", ""),
            ])

        output = BytesIO()
        wb.save(output)
        output.seek(0)

        return send_file(
            output,
            as_attachment=True,
            download_name=f"{filename_base}.xlsx",
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

    if fmt == "sql":
        lines = []

        for c in clients:
            name = str(c.get("name", "")).replace("'", "\\'")
            email = str(c.get("email", "")).replace("'", "\\'")
            phone = str(c.get("phone", "")).replace("'", "\\'")
            address = str(c.get("address", "")).replace("'", "\\'")
            notes = str(c.get("notes", "")).replace("'", "\\'")

            lines.append(
                "INSERT INTO clients (name,email,phone,address,notes) "
                f"VALUES ('{name}','{email}','{phone}','{address}','{notes}');"
            )

        return Response(
            "\n".join(lines),
            mimetype="text/plain",
            headers={
                "Content-Disposition": f"attachment; filename={filename_base}.sql"
            }
        )

    if fmt == "pdf":
        output = BytesIO()
        pdf = canvas.Canvas(output, pagesize=letter)
        width, height = letter

        y = height - 50

        pdf.setFont("Helvetica-Bold", 14)
        pdf.drawString(40, y, "Backup de Clientes")
        y -= 20

        pdf.setFont("Helvetica", 9)
        pdf.drawString(40, y, f"Empresa: {backup.get('company_name', '')}")
        y -= 15
        pdf.drawString(40, y, f"Total clientes: {backup.get('total_clients', 0)}")
        y -= 30

        pdf.setFont("Helvetica-Bold", 8)
        pdf.drawString(40, y, "Nombre | Correo | Teléfono | Dirección | Notas")
        y -= 15

        pdf.setFont("Helvetica", 8)

        for c in clients:
            line = (
                f"{c.get('name', '')} | "
                f"{c.get('email', '')} | "
                f"{c.get('phone', '')} | "
                f"{c.get('address', '')} | "
                f"{c.get('notes', '')}"
            )

            if len(line) > 120:
                line = line[:117] + "..."

            pdf.drawString(40, y, line)
            y -= 14

            if y < 50:
                pdf.showPage()
                y = height - 50
                pdf.setFont("Helvetica", 8)

        pdf.save()
        output.seek(0)

        return send_file(
            output,
            as_attachment=True,
            download_name=f"{filename_base}.pdf",
            mimetype="application/pdf"
        )

    flash("Formato no válido.", "warning")
    return redirect(url_for("clients.clients_backups"))